"""Stage 5.1 data-quality repair for Stage 5 BUSINESS_DATA workbooks.

Fixes eight issue classes against as_of_date=2026-09-12 without touching
Canonical facts, Stage 3/4 assets, templates or repository structure:
1. time consistency (future hires, future completions, premature revenue)
2. org reporting lines (intern/P1 department roots rebuilt)
3. project-manager qualification for P004-P012
4. DATA011 skill tags rebuilt by role/department pools
5. timesheet self-approval removed
6. project delay status + supplier recheck realism
7. DATA008 -> DATA006 procurement cost linkage
8. DATA012 metric dictionary formula repair
"""
from __future__ import annotations

import hashlib
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = ROOT / "knowledge" / "business_data"
AS_OF = date(2026, 9, 12)
DATE_FMT = "yyyy-mm-dd"


def wb_path(dsid: str) -> Path:
    return next(DATA_DIR.glob(f"{dsid}_*.xlsx"))


def sheet_map(ws):
    headers = [c.value for c in ws[1]]
    return {str(h): i + 1 for i, h in enumerate(headers)}


def rows_by_key(ws, cols, key):
    out = {}
    for r in range(2, ws.max_row + 1):
        out[str(ws.cell(r, cols[key]).value)] = r
    return out


def set_date(ws, row, col, value):
    cell = ws.cell(row, col)
    if value is None:
        cell.value = None
    else:
        cell.value = datetime(value.year, value.month, value.day)
        cell.number_format = DATE_FMT


def set_value(ws, row, col, value):
    ws.cell(row, col).value = value


def as_date(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return date.fromisoformat(str(v))


# ---------------------------------------------------------------- DATA001
def fix_employees():
    wb = openpyxl.load_workbook(wb_path("DATA001"))
    ws = wb["数据"]
    c = sheet_map(ws)
    rows = rows_by_key(ws, c, "employee_id")

    # 1a. future hires moved one year into the past
    future_hires = 0
    for r in rows.values():
        hd = as_date(ws.cell(r, c["hire_date"]).value)
        if hd and hd > AS_OF:
            set_date(ws, r, c["hire_date"], hd.replace(year=hd.year - 1))
            future_hires += 1

    # 1b. intern mis-graded at manager level
    set_value(ws, rows["EMP0072"], c["job_level"], "P1")

    # 2. department heads (position -> director) and team leads
    heads = {
        "D01": ("EMP0018", "市场总监"), "D02": ("EMP0102", "研发总监"),
        "D03": ("EMP0204", "交付总监"), "D04": ("EMP0219", "采购总监"),
        "D05": ("EMP0242", "人力资源总监"), "D06": ("EMP0252", "财务总监"),
        "D07": ("EMP0268", "法务总监"), "D08": ("EMP0285", "行政总监"),
    }
    leads = {
        "D01": ["EMP0013", "EMP0008", "EMP0033"],
        "D02": ["EMP0042", "EMP0057", "EMP0062", "EMP0067", "EMP0092", "EMP0097", "EMP0127"],
        "D03": ["EMP0001", "EMP0002", "EMP0003", "EMP0139", "EMP0169", "EMP0174"],
        "D04": ["EMP0209", "EMP0214"],
        "D05": ["EMP0232", "EMP0237"],
        "D06": ["EMP0247", "EMP0257"],
        "D07": ["EMP0263"],
        "D08": ["EMP0275", "EMP0280"],
    }
    members = {}
    for r in rows.values():
        did = str(ws.cell(r, c["department_id"]).value)
        eid = str(ws.cell(r, c["employee_id"]).value)
        members.setdefault(did, []).append(eid)
    for did, member_ids in members.items():
        head_id, head_title = heads[did]
        set_value(ws, rows[head_id], c["position_name"], head_title)
        set_value(ws, rows[head_id], c["manager_employee_id"], None)
        lead_ids = leads[did]
        for lid in lead_ids:
            set_value(ws, rows[lid], c["manager_employee_id"], head_id)
        targets = lead_ids if len(lead_ids) > 1 else [head_id] + lead_ids
        others = sorted(m for m in member_ids if m != head_id and m not in lead_ids)
        for i, mid in enumerate(others):
            set_value(ws, rows[mid], c["manager_employee_id"], targets[i % len(targets)])

    # 3. project-manager promotions (P1 -> P3) for P004/P008/P009/P010 managers
    for eid in ["EMP0150", "EMP0145", "EMP0180", "EMP0175"]:
        set_value(ws, rows[eid], c["job_level"], "P3")

    wb.save(wb_path("DATA001"))
    return {"future_hires_shifted": future_hires}


# ---------------------------------------------------------------- DATA003
def fix_suppliers():
    wb = openpyxl.load_workbook(wb_path("DATA003"))
    ws = wb["数据"]
    c = sheet_map(ws)
    rows = rows_by_key(ws, c, "supplier_id")
    renewed = {"SUP003": date(2028, 3, 15), "SUP007": date(2028, 7, 15), "SUP013": date(2028, 1, 15),
               "SUP015": date(2028, 3, 15), "SUP019": date(2028, 7, 15), "SUP025": date(2028, 1, 15)}
    for sid, due in renewed.items():
        set_date(ws, rows[sid], c["recheck_due_date"], due)
    for sid in ["SUP005", "SUP017", "SUP029"]:
        set_value(ws, rows[sid], c["status"], "观察")
        set_value(ws, rows[sid], c["qualification_status"], "附条件通过")
    # already under observation; align qualification with conditional-use state
    set_value(ws, rows["SUP027"], c["qualification_status"], "附条件通过")
    wb.save(wb_path("DATA003"))
    return {"suppliers_renewed": len(renewed), "suppliers_conditional": 3}


# ------------------------------------------------------- DATA004 projects
NEW_PM = {  # project -> (employee_id, name); P008 keeps EMP0145 (promoted)
    "P004": ("EMP0150", "合成员工150"),
    "P009": ("EMP0180", "合成员工180"),
    "P010": ("EMP0175", "合成员工175"),
}
DELAYED = {  # project -> (forecast_end, new_risk or None)
    "P005": (date(2026, 12, 31), "L3"),
    "P006": (date(2026, 12, 31), "L3"),
    "P007": (date(2026, 11, 30), "L2"),
    "P008": (date(2026, 11, 30), None),
    "P009": (date(2026, 10, 31), None),
}
RECOGNIZED_FIX = {"P005": 1344000, "P006": 1890000, "P007": 705600, "P008": 1092000,
                  "P009": 274400, "P010": 630000, "P011": 728000, "P012": 294000}


def fix_projects():
    path = wb_path("DATA004")
    wb = openpyxl.load_workbook(path)
    ws = wb["数据"]
    c = sheet_map(ws)
    rows = rows_by_key(ws, c, "project_id")

    # add forecast_end_date column after actual_end_date
    insert_at = c["actual_end_date"] + 1
    ws.insert_cols(insert_at)
    ws.cell(1, insert_at).value = "forecast_end_date"
    c = sheet_map(ws)
    for pid, r in rows.items():
        forecast = DELAYED.get(pid, (None, None))[0]
        set_date(ws, r, insert_at, forecast)

    for pid, (forecast, risk) in DELAYED.items():
        set_value(ws, rows[pid], c["project_status"], "delayed")
        if risk:
            set_value(ws, rows[pid], c["highest_risk_level"], risk)
    for pid, amount in RECOGNIZED_FIX.items():
        set_value(ws, rows[pid], c["recognized_revenue"], amount)
    for pid, (eid, name) in NEW_PM.items():
        set_value(ws, rows[pid], c["project_manager_employee_id"], eid)
        set_value(ws, rows[pid], c["project_manager_name"], name)

    # field dictionary: insert matching row
    ds = wb["字段字典"]
    dict_row = c["actual_end_date"] + 2  # header + position of new field
    ds.insert_rows(dict_row)
    values = ["forecast_end_date", "date", "预测完成日期；延期项目填最新预测，按期或已完成项目留空",
              None, None, "按业务关系", "INTERNAL"]
    for i, v in enumerate(values, start=1):
        ds.cell(dict_row, i).value = v
    wb.save(path)
    return {"delayed_projects": len(DELAYED), "pm_reassigned": len(NEW_PM)}


# ------------------------------------------------------- DATA005 contracts
def fix_contracts():
    path = wb_path("DATA005")
    wb = openpyxl.load_workbook(path)
    ws = wb["数据"]
    c = sheet_map(ws)
    rows = rows_by_key(ws, c, "contract_id")

    # purchase-contract realignment, mirrors DATA006 procurement windows
    completed = {
        "PC005": (date(2025, 6, 10), date(2025, 6, 28)),
        "PC007": (date(2026, 1, 10), date(2026, 1, 28)),
        "PC009": (date(2026, 7, 10), date(2026, 7, 28)),
        "PC010": (date(2026, 3, 10), date(2026, 3, 28)),
        "PC011": (date(2026, 8, 10), date(2026, 8, 28)),
        "PC013": (date(2026, 7, 10), date(2026, 7, 28)),
        "PC020": (date(2026, 3, 10), date(2026, 3, 28)),
        "PC022": (date(2026, 7, 15), date(2026, 8, 5)),
        "PC023": (date(2025, 8, 10), date(2025, 8, 28)),
        "PC026": (date(2025, 8, 10), date(2025, 8, 28)),
        "PC028": (date(2026, 2, 10), date(2026, 2, 28)),
        "PC032": (date(2025, 9, 10), date(2025, 9, 24)),
    }
    for cid, (sign, end) in completed.items():
        r = rows[cid]
        set_date(ws, r, c["sign_date"], sign)
        set_date(ws, r, c["start_date"], sign)
        set_date(ws, r, c["end_date"], end)
        set_value(ws, r, c["contract_status"], "completed")
    activated = {
        "PC012": date(2026, 5, 15), "PC018": date(2025, 10, 15), "PC021": date(2026, 8, 25),
        "PC024": date(2025, 4, 15), "PC030": date(2026, 6, 15), "PC033": date(2026, 8, 20),
        "PC035": date(2026, 8, 15), "PC036": date(2025, 11, 15),
    }
    for cid, sign in activated.items():
        r = rows[cid]
        set_date(ws, r, c["sign_date"], sign)
        set_date(ws, r, c["start_date"], sign)
        set_date(ws, r, c["end_date"], None)
        set_value(ws, r, c["contract_status"], "active")

    # sales-contract owner follows the reassigned project managers
    owners = {"P004-C01": "EMP0150", "P009-C01": "EMP0180", "P010-C01": "EMP0175"}
    for cid, eid in owners.items():
        set_value(ws, rows[cid], c["owner_employee_id"], eid)
    wb.save(path)
    return {"purchase_contracts_realigned": len(completed) + len(activated)}


# ------------------------------------------------------ DATA006 procurement
# procurement -> (request, delivery|None, status, committed)
PROC_FIX = {
    "PR-P001-01": (date(2024, 3, 20), date(2024, 6, 10), "completed", date(2024, 6, 10)),
    "PR-P002-01": (date(2024, 7, 22), date(2024, 9, 2), "delayed", date(2024, 8, 15)),
    "PR-P003-01": (date(2024, 5, 15), date(2024, 9, 2), "completed", date(2024, 8, 20)),
    "PR-P003-02": (date(2024, 6, 20), date(2024, 10, 8), "completed", date(2024, 10, 8)),
    "PR005": (date(2025, 6, 10), date(2025, 6, 28), "completed", date(2025, 6, 28)),
    "PR006": (date(2025, 7, 10), None, "approved", date(2026, 10, 31)),
    "PR007": (date(2026, 1, 10), date(2026, 1, 28), "completed", date(2026, 1, 28)),
    "PR008": (date(2025, 9, 10), date(2025, 9, 28), "completed", date(2025, 9, 20)),
    "PR009": (date(2026, 7, 10), date(2026, 7, 28), "completed", date(2026, 7, 28)),
    "PR010": (date(2026, 3, 10), date(2026, 3, 28), "completed", date(2026, 3, 28)),
    "PR011": (date(2026, 8, 10), date(2026, 8, 28), "completed", date(2026, 8, 28)),
    "PR012": (date(2026, 5, 15), None, "approved", date(2026, 12, 31)),
    "PR013": (date(2026, 7, 10), date(2026, 7, 28), "completed", date(2026, 7, 28)),
    "PR014": (date(2025, 3, 10), date(2025, 3, 28), "completed", date(2025, 3, 28)),
    "PR015": (date(2026, 4, 10), date(2026, 4, 28), "completed", date(2026, 4, 20)),
    "PR016": (date(2025, 5, 10), date(2025, 5, 28), "completed", date(2025, 5, 28)),
    "PR017": (date(2026, 6, 10), date(2026, 6, 28), "completed", date(2026, 6, 28)),
    "PR018": (date(2025, 10, 15), None, "approved", date(2026, 11, 15)),
    "PR019": (date(2026, 8, 10), date(2026, 8, 28), "completed", date(2026, 8, 20)),
    "PR020": (date(2026, 3, 10), date(2026, 3, 28), "completed", date(2026, 3, 28)),
    "PR021": (date(2026, 8, 25), None, "approved", date(2026, 10, 28)),
    "PR022": (date(2026, 7, 15), date(2026, 8, 5), "completed", date(2026, 8, 5)),
    "PR023": (date(2025, 8, 10), date(2025, 8, 28), "completed", date(2025, 8, 28)),
    "PR024": (date(2025, 4, 15), None, "approved", date(2026, 11, 30)),
    "PR025": (date(2026, 2, 10), date(2026, 2, 28), "completed", date(2026, 2, 28)),
    "PR026": (date(2025, 8, 10), date(2025, 8, 28), "completed", date(2025, 8, 28)),
    "PR027": (date(2026, 4, 10), date(2026, 4, 28), "completed", date(2026, 4, 28)),
    "PR028": (date(2026, 2, 10), date(2026, 2, 28), "completed", date(2026, 2, 28)),
    "PR029": (date(2026, 6, 10), date(2026, 6, 28), "completed", date(2026, 6, 28)),
    "PR030": (date(2026, 6, 15), None, "approved", date(2027, 1, 31)),
    "PR031": (date(2026, 8, 10), date(2026, 8, 28), "completed", date(2026, 8, 28)),
    "PR032": (date(2025, 9, 10), date(2025, 9, 24), "completed", date(2025, 9, 24)),
    "PR033": (date(2026, 8, 20), None, "approved", date(2026, 11, 30)),
    "PR034": (date(2025, 11, 10), date(2025, 11, 28), "completed", date(2025, 11, 28)),
    "PR035": (date(2026, 8, 15), None, "approved", date(2026, 11, 15)),
    "PR036": (date(2025, 11, 15), None, "approved", date(2026, 10, 31)),
}


def fix_procurement():
    path = wb_path("DATA006")
    wb = openpyxl.load_workbook(path)
    ws = wb["数据"]
    c = sheet_map(ws)
    rows = rows_by_key(ws, c, "procurement_id")

    insert_at = c["delivery_date"] + 1
    ws.insert_cols(insert_at)
    ws.cell(1, insert_at).value = "committed_delivery_date"
    c = sheet_map(ws)
    for prid, (req, dlv, status, committed) in PROC_FIX.items():
        r = rows[prid]
        set_date(ws, r, c["request_date"], req)
        set_date(ws, r, c["delivery_date"], dlv)
        set_value(ws, r, c["status"], status)
        set_date(ws, r, c["committed_delivery_date"], committed)

    ds = wb["字段字典"]
    dict_row = c["delivery_date"] + 2
    ds.insert_rows(dict_row)
    values = ["committed_delivery_date", "date", "承诺交付日期；供应商承诺的交付截止日",
              None, None, "按业务关系", "SENSITIVE"]
    for i, v in enumerate(values, start=1):
        ds.cell(dict_row, i).value = v
    wb.save(path)
    approved = sum(1 for v in PROC_FIX.values() if v[2] == "approved")
    return {"procurement_realigned": len(PROC_FIX), "procurement_approved_open": approved}


# ------------------------------------------------------ DATA007 timesheets
def fix_timesheets():
    path = wb_path("DATA007")
    wb = openpyxl.load_workbook(path)
    ws = wb["数据"]
    c = sheet_map(ws)
    pm = {"P001": "EMP0001", "P002": "EMP0002", "P003": "EMP0003", "P004": "EMP0150",
          "P005": "EMP0142", "P006": "EMP0143", "P007": "EMP0144", "P008": "EMP0145",
          "P009": "EMP0180", "P010": "EMP0175", "P011": "EMP0148", "P012": "EMP0149"}
    dept_head = "EMP0204"
    changed = 0
    for r in range(2, ws.max_row + 1):
        pid = str(ws.cell(r, c["project_id"]).value)
        eid = str(ws.cell(r, c["employee_id"]).value)
        approver = dept_head if eid == pm[pid] else pm[pid]
        if str(ws.cell(r, c["approver_employee_id"]).value) != approver:
            ws.cell(r, c["approver_employee_id"]).value = approver
            changed += 1
    wb.save(path)
    return {"timesheet_approvers_rebuilt": changed}


# ------------------------------------------------------------ DATA008 costs
COST_LINKS = {
    "COST00002": "PR-P001-01", "COST00010": "PR-P002-01",
    "COST00017": "PR-P003-01", "COST00018": "PR-P003-02",
    "COST00025": "PR032", "COST00026": "PR023",
    "COST00031": "PR024", "COST00032": "PR015",
    "COST00037": "PR025", "COST00038": "PR007",
    "COST00043": "PR017", "COST00044": "PR026",
    "COST00049": "PR009", "COST00050": "PR027", "COST00052": "PR036",
    "COST00056": "PR010", "COST00058": "PR028",
    "COST00061": "PR029", "COST00062": "PR011", "COST00064": "PR020",
    "COST00067": "PR021", "COST00070": "PR012",
    "COST00073": "PR022", "COST00074": "PR031",
}


def fix_costs():
    path = wb_path("DATA008")
    wb = openpyxl.load_workbook(path)
    ws = wb["数据"]
    c = sheet_map(ws)
    rows = rows_by_key(ws, c, "cost_id")
    for cid, prid in COST_LINKS.items():
        set_value(ws, rows[cid], c["related_procurement_id"], prid)
    wb.save(path)
    return {"cost_procurement_links": len(COST_LINKS)}


# ---------------------------------------------------------- DATA009 budget
def fix_budgets():
    path = wb_path("DATA009")
    wb = openpyxl.load_workbook(path)
    ws = wb["数据"]
    c = sheet_map(ws)
    swap = {"EMP0141": "EMP0150", "EMP0146": "EMP0180", "EMP0147": "EMP0175"}
    changed = 0
    for r in range(2, ws.max_row + 1):
        owner = str(ws.cell(r, c["owner_employee_id"]).value)
        if owner in swap:
            ws.cell(r, c["owner_employee_id"]).value = swap[owner]
            changed += 1
    wb.save(path)
    return {"budget_owners_swapped": changed}


# ---------------------------------------------------------- DATA010 revenue
def fix_revenue():
    path = wb_path("DATA010")
    wb = openpyxl.load_workbook(path)
    ws = wb["数据"]
    c = sheet_map(ws)
    rows = rows_by_key(ws, c, "revenue_id")
    unearned = ["REV00017", "REV00020", "REV00023", "REV00026", "REV00029",
                "REV00032", "REV00035", "REV00038"]
    for rid in unearned:
        r = rows[rid]
        set_date(ws, r, c["recognized_date"], None)
        set_value(ws, r, c["billed_amount"], 0)
        set_value(ws, r, c["recognized_amount"], 0)
        set_value(ws, r, c["outstanding_amount"], 0)
    # P004 milestone recognized after project closure moved into execution window
    r = rows["REV00013"]
    set_date(ws, r, c["due_date"], date(2025, 6, 30))
    set_date(ws, r, c["recognized_date"], date(2025, 6, 30))
    wb.save(path)
    return {"premature_revenue_cleared": len(unearned)}


# ------------------------------------------------------------ DATA011 skills
SKILL_CODE = {"项目管理": "SK001", "需求分析": "SK002", "解决方案设计": "SK003", "Java开发": "SK004",
              "Python开发": "SK005", "算法建模": "SK006", "知识工程": "SK007", "数据治理": "SK008",
              "测试管理": "SK009", "实施交付": "SK010", "采购管理": "SK011", "合同审查": "SK012",
              "预算分析": "SK013", "招聘面试": "SK014", "培训设计": "SK015", "文档规范": "SK016"}
DEPT_POOL = {
    "D01": ["解决方案设计", "需求分析", "项目管理", "文档规范"],
    "D02": ["Java开发", "Python开发", "算法建模", "知识工程", "数据治理", "测试管理", "需求分析", "解决方案设计", "项目管理", "文档规范"],
    "D03": ["项目管理", "需求分析", "实施交付", "解决方案设计", "数据治理", "测试管理", "知识工程", "文档规范"],
    "D04": ["采购管理", "合同审查", "预算分析", "项目管理", "文档规范"],
    "D05": ["招聘面试", "培训设计", "项目管理", "文档规范"],
    "D06": ["预算分析", "合同审查", "项目管理", "文档规范"],
    "D07": ["合同审查", "数据治理", "项目管理", "文档规范"],
    "D08": ["文档规范", "培训设计", "采购管理", "预算分析", "项目管理"],
}
PRIMARY_SKILL = {
    "解决方案顾问": "解决方案设计", "客户经理": "需求分析", "投标专员": "解决方案设计", "市场运营": "需求分析",
    "软件工程师": "Java开发", "算法工程师": "算法建模", "架构师": "解决方案设计", "产品经理": "需求分析",
    "测试工程师": "测试管理", "知识工程师": "知识工程",
    "项目经理": "项目管理", "实施顾问": "实施交付", "交付工程师": "实施交付", "质量经理": "测试管理",
    "数据治理工程师": "数据治理",
    "采购专员": "采购管理", "供应商管理员": "合同审查", "采购经理": "采购管理",
    "HRBP": "招聘面试", "招聘专员": "招聘面试", "培训专员": "培训设计", "绩效专员": "培训设计",
    "财务经理": "预算分析", "项目会计": "预算分析", "预算分析师": "预算分析", "应收会计": "预算分析",
    "法务经理": "合同审查", "合同管理员": "合同审查", "风控专员": "合同审查", "数据合规专员": "数据治理",
    "综合行政专员": "文档规范", "档案管理员": "文档规范", "会议运营": "文档规范", "行政经理": "文档规范",
    "市场总监": "项目管理", "研发总监": "项目管理", "交付总监": "项目管理", "采购总监": "项目管理",
    "人力资源总监": "项目管理", "财务总监": "项目管理", "法务总监": "项目管理", "行政总监": "项目管理",
}
PM_WITH_SKILL = {"EMP0001", "EMP0002", "EMP0003", "EMP0150", "EMP0142", "EMP0143", "EMP0144",
                 "EMP0145", "EMP0180", "EMP0175", "EMP0148", "EMP0149"}
PRIMARY_PROF = {"M1": "L4-指导", "P4": "L4-指导", "P3": "L3-独立", "P2": "L3-独立", "P1": "L2-可协作"}
STEP_DOWN = {"L4-指导": "L3-独立", "L3-独立": "L2-可协作", "L2-可协作": "L1-了解", "L1-了解": "L1-了解"}


def _hash(*parts: str) -> int:
    return int(hashlib.md5("|".join(parts).encode("utf-8")).hexdigest(), 16)


def rebuild_skills():
    emp_wb = openpyxl.load_workbook(wb_path("DATA001"))
    emp_ws = emp_wb["数据"]
    ec = sheet_map(emp_ws)
    employees = []
    for r in range(2, emp_ws.max_row + 1):
        employees.append({
            "employee_id": str(emp_ws.cell(r, ec["employee_id"]).value),
            "department_id": str(emp_ws.cell(r, ec["department_id"]).value),
            "position_name": str(emp_ws.cell(r, ec["position_name"]).value),
            "job_level": str(emp_ws.cell(r, ec["job_level"]).value),
            "employment_type": str(emp_ws.cell(r, ec["employment_type"]).value),
            "manager_employee_id": emp_ws.cell(r, ec["manager_employee_id"]).value,
        })

    path = wb_path("DATA011")
    wb = openpyxl.load_workbook(path)
    ws = wb["数据"]
    c = sheet_map(ws)
    old_source = {}
    for r in range(2, ws.max_row + 1):
        old_source[str(ws.cell(r, c["skill_tag_id"]).value)] = str(ws.cell(r, c["source_type"]).value)

    tag_row = 2
    for emp in employees:
        eid = emp["employee_id"]
        pool = DEPT_POOL[emp["department_id"]]
        primary = PRIMARY_SKILL.get(emp["position_name"], pool[0])
        rest = [s for s in pool if s != primary]
        secondary = rest[_hash(eid, "sec") % len(rest)]
        if emp["position_name"] in {"软件工程师", "算法工程师"} and "Python开发" in rest and _hash(eid, "py") % 3:
            secondary = "Python开发"
        if eid in PM_WITH_SKILL and primary != "项目管理":
            secondary = "项目管理"
        third_pool = [s for s in rest if s != secondary]
        third = third_pool[_hash(eid, "3rd") % len(third_pool)]
        if emp["employment_type"] == "实习":
            prof1 = "L1-了解"
        else:
            prof1 = PRIMARY_PROF[emp["job_level"]]
        prof2 = STEP_DOWN[prof1]
        prof3 = "L2-可协作" if _hash(eid, "prof3") % 2 else "L1-了解"
        reviewer = emp["manager_employee_id"] or "EMP0242"
        if eid == "EMP0242":
            reviewer = "EMP0204"
        for skill, prof, evid in ((primary, prof1, "岗位验证"), (secondary, prof2, "项目交付"), (third, prof3, "培训考核")):
            tag_id = f"TAG{tag_row - 1:06d}"
            verified = date(2025, 1, 15) + timedelta(days=_hash(eid, skill) % 560)
            ws.cell(tag_row, c["skill_tag_id"]).value = tag_id
            ws.cell(tag_row, c["employee_id"]).value = eid
            ws.cell(tag_row, c["skill_code"]).value = SKILL_CODE[skill]
            ws.cell(tag_row, c["skill_name"]).value = skill
            ws.cell(tag_row, c["proficiency_level"]).value = prof
            ws.cell(tag_row, c["evidence_type"]).value = evid
            set_date(ws, tag_row, c["last_verified_date"], verified)
            ws.cell(tag_row, c["reviewer_employee_id"]).value = reviewer
            ws.cell(tag_row, c["tag_status"]).value = "有效"
            ws.cell(tag_row, c["source_type"]).value = old_source.get(tag_id, "SYNTHETIC")
            tag_row += 1
    assert tag_row - 2 == 858, tag_row
    wb.save(path)
    return {"skill_tags_rebuilt": 858}


# ------------------------------------------------------------ DATA012 metrics
METRIC_FIX = {
    "OP005": {"calculation_formula": "SUM(amount WHERE accounting_status='posted')"},
    "OP009": {"source_dataset": "项目主数据", "calculation_rule": "收入为零时不可用"},
    "OP010": {"source_dataset": "项目主数据", "calculation_rule": "合同额为零时不可用"},
    "OP011": {"calculation_formula": "COUNT(employment_status='在职')"},
    "OP014": {"calculation_formula": "SUM(work_hours WHERE approval_status='approved')"},
    "OP016": {"calculation_formula": "SUM(amount WHERE status IN 'approved','completed','delayed')"},
    "OP017": {"calculation_formula": "COUNT(delivery_date<=committed_delivery_date AND status IN 'completed','delayed')/COUNT(status IN 'completed','delayed')",
              "calculation_rule": "未交付订单不进入分母；按期以承诺交付日期为准"},
    "OP019": {"calculation_formula": "COUNT(risk_level IN 'L3','L4')"},
    "OP020": {"calculation_formula": "COUNT(contract_status IN 'disputed','overdue')"},
    "OP021": {"calculation_formula": "COUNT(project_status IN 'completed','completed_with_issues' AND actual_end_date<=planned_end_date)/COUNT(project_status IN 'completed','completed_with_issues')"},
    "OP022": {"calculation_formula": "MAX(0, actual_end_date-planned_end_date)"},
    "OP024": {"definition": "进行中或延期项目数",
              "calculation_formula": "COUNT(project_status IN 'in_progress','delayed')"},
    "OP028": {"calculation_formula": "COUNT(amount>=50000 AND method IN '三家比选','竞争性采购','公开招标或专项审批')/COUNT(amount>=50000)",
              "calculation_rule": "按内部阈值，金额5万及以上应采用竞争方式"},
    "OP029": {"calculation_formula": "COUNT_DISTINCT(人员能力标签.employee_id WHERE tag_status='有效')/COUNT(员工花名册.employment_status='在职')"},
    "OP030": {"calculation_formula": "COUNT_DISTINCT(employee_id WHERE proficiency_level='L4-指导')"},
    "OP031": {"calculation_formula": "SUM(approved_amount WHERE approval_status='approved')"},
    "OP033": {"calculation_formula": "COUNT(purchase_contract_id NOT NULL)/COUNT(procurement_id)",
              "calculation_rule": "按purchase_contract_id关联合同台账contract_id"},
    "OP034": {"calculation_formula": "N/A", "status": "draft",
              "calculation_rule": "无法由单一数据集字段公式计算，由质量门校验 scripts/validate_all.py 输出"},
    "OP035": {"calculation_formula": "N/A", "status": "draft",
              "calculation_rule": "无法由单一数据集字段公式计算，由质量门外键校验输出"},
}


def fix_metrics():
    path = wb_path("DATA012")
    wb = openpyxl.load_workbook(path)
    ws = wb["数据"]
    c = sheet_map(ws)
    rows = rows_by_key(ws, c, "metric_id")
    for mid, changes in METRIC_FIX.items():
        for field, value in changes.items():
            set_value(ws, rows[mid], c[field], value)
    wb.save(path)
    return {"metrics_fixed": len(METRIC_FIX)}


def main():
    summary = {}
    summary.update(fix_employees())
    summary.update(fix_suppliers())
    summary.update(fix_projects())
    summary.update(fix_contracts())
    summary.update(fix_procurement())
    summary.update(fix_timesheets())
    summary.update(fix_costs())
    summary.update(fix_budgets())
    summary.update(fix_revenue())
    summary.update(rebuild_skills())
    summary.update(fix_metrics())
    for key, value in summary.items():
        print(f"{key}: {value}")


if __name__ == "__main__":
    main()
